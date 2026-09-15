import base64
import io
import pytest
import pypdf
import docx

from app.utils.attachment_parser import (
    decode_data_url,
    extract_pdf_text,
    extract_docx_text,
    extract_text_content,
    format_message_content,
    process_attachment,
)


def test_decode_data_url():
    sample_text = "Hello world attachment"
    b64 = base64.b64encode(sample_text.encode("utf-8")).decode("utf-8")
    data_url = f"data:text/plain;base64,{b64}"

    mime, raw = decode_data_url(data_url)
    assert mime == "text/plain"
    assert raw.decode("utf-8") == sample_text


def test_extract_text_content():
    raw = "def hello():\n    return 'world'".encode("utf-8")
    result = extract_text_content(raw)
    assert "def hello():" in result


def test_extract_pdf_text():
    # Create an in-memory PDF
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = io.BytesIO()
    writer.write(stream)
    pdf_bytes = stream.getvalue()

    result = extract_pdf_text(pdf_bytes)
    assert isinstance(result, str)


def test_extract_docx_text():
    doc = docx.Document()
    doc.add_paragraph("Paragraph inside docx")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Header 1"
    table.rows[0].cells[1].text = "Header 2"

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    result = extract_docx_text(docx_bytes)
    assert "Paragraph inside docx" in result
    assert "Header 1 | Header 2" in result


def test_format_message_content_plain_text_only():
    content = "Hello there"
    formatted = format_message_content(content, None)
    assert formatted == "Hello there"


def test_format_message_content_with_document():
    raw = "Content of notes.txt".encode("utf-8")
    b64 = base64.b64encode(raw).decode("utf-8")
    attachments = [
        {
            "name": "notes.txt",
            "size": len(raw),
            "type": "text/plain",
            "dataUrl": f"data:text/plain;base64,{b64}",
        }
    ]

    formatted = format_message_content("Summarize this", attachments)
    assert isinstance(formatted, str)
    assert "[Attached Document: notes.txt" in formatted
    assert "Content of notes.txt" in formatted
    assert "Summarize this" in formatted


def test_format_message_content_with_image():
    attachments = [
        {
            "name": "photo.png",
            "size": 100,
            "type": "image/png",
            "dataUrl": "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=",
        }
    ]

    formatted = format_message_content("What is in this image?", attachments)
    assert isinstance(formatted, list)
    assert formatted[0] == {"type": "text", "text": "What is in this image?"}
    assert formatted[1]["type"] == "image_url"
    assert formatted[1]["image_url"]["url"].startswith("data:image/png;base64,")
